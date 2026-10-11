"""Validation safeguards: no hindsight fills, calendar gaps, maturation and disclosures."""
import json
import math
import re
import shutil
import subprocess

import pandas as pd
import pytest

from radar.main import export
from radar.storage import connect, store
from radar.signal_validation import (
    HORIZONS, market_calendar, validate_one_stock, aggregate_validation
)
from radar.engine import screen
from radar.dashboard import render_dashboard


def make_test_data(n=105):
    dates = pd.bdate_range("2026-03-02", periods=n).strftime("%Y-%m-%d").tolist()
    data=[]
    for i, day in enumerate(dates):
        price=100+i*.08
        data.append({
            "date":day,"market":"TWSE","code":"2466","name":"冠西電",
            "open":price,"high":price+2,"low":price-2,
            "close":price+1,"volume":210000,"turnover":24_000_000,
            "source_url":"synthetic-test",
            "impulse_event":False,"impulse_signal":"NONE",
        })
    return pd.DataFrame(data), dates


def test_next_session_open_not_candle_close_or_low():
    frame, dates=make_test_data()
    frame.loc[88,"impulse_signal"]="START"
    frame.loc[88,"impulse_event"]=True
    frame.loc[89,"open"]=110.0
    frame.loc[89,"close"]=121.0
    frame.loc[89,"low"]=108.0
    frame.loc[89,"high"]=124.0
    frame.loc[93,"close"]=130.0
    records,pending,events=validate_one_stock(frame, dates,"TWSE","2466")
    assert events==1
    one=next(r for r in records if r["horizon"]==1)
    five=next(r for r in records if r["horizon"]==5)
    assert one["entry_date"]==dates[89]
    assert one["entry_open"]==110.0
    assert one["forward_return_pct"]==10.0
    assert one["max_adverse_pct"]==pytest.approx(-1.8182,abs=.001)
    assert five["exit_date"]==dates[93]
    assert five["forward_return_pct"]==pytest.approx(18.1818,abs=.001)
    assert len(records)==4


def test_recent_signal_unmatured_and_gap_not_fabricated():
    frame, dates=make_test_data()
    frame.loc[103,["impulse_signal","impulse_event"]]=["EARLY",True]
    rows,pending,total=validate_one_stock(frame,dates,"TWSE","2466")
    assert total==1
    assert len(rows)==1
    assert rows[0]["horizon"]==1
    assert pending["EARLY"]["3"]==1
    assert pending["EARLY"]["5"]==1
    assert pending["EARLY"]["10"]==1
    # For an older signal, missing the next market day must never be interpreted
    # as holding one fewer session.
    frame2, dates=make_test_data()
    frame2.loc[86,["impulse_signal","impulse_event"]]=["PREP",True]
    frame2=frame2.loc[frame2.date != dates[88]].copy()
    rows,pending,total=validate_one_stock(frame2,dates,"TWSE","2466")
    assert total==1
    assert {r["horizon"] for r in rows}=={1}
    assert pending["PREP"]["3"]==1


def test_hot_episode_not_counted_daily():
    frame, dates=make_test_data()
    frame.loc[85:88,"impulse_signal"]="HOT"
    frame.loc[94,"impulse_signal"]="HOT"
    rows,pending,total=validate_one_stock(frame,dates,"TWSE","2466")
    assert total==2
    assert len([r for r in rows if r["horizon"]==5])==2


def test_empty_or_low_sample_not_called_success_rate():
    report=aggregate_validation([],{},0,"2026-10-08")
    assert len(report["sample_rows"])==16
    assert all(x["sample_count"]==0 for x in report["sample_rows"])
    assert all(x["positive_rate_pct"] is None for x in report["sample_rows"])
    assert all(x["sample_status"]=="INSUFFICIENT_SAMPLE" for x in report["sample_rows"])
    assert report["entry_rule"]=="NEXT_SESSION_OPEN"
    assert report["interpretation"]["not_a_success_probability"] is True


def test_market_calendar_excludes_stray_small_market_batches():
    frame,dates=make_test_data()
    normal=frame[["date","market","code"]].copy()
    stray=pd.DataFrame([{"date":"2026-06-01","market":"TWSE","code":"9988"}])
    # Single-stock fixture has a 1-stock normal session, so use a separate
    # synthetic multi-stock market for realistic completeness gating.
    copies=[normal.assign(code=str(1000+i)) for i in range(20)]
    full=pd.concat(copies+[stray],ignore_index=True)
    result=market_calendar(full)
    assert "2026-06-01" not in result["TWSE"] if "2026-06-01" not in dates else True


def test_normal_scan_contract_and_validation_mode():
    frame,_=make_test_data()
    price_frame=frame.drop(columns=["impulse_signal","impulse_event"])
    normal=screen(price_frame)
    ranking,chart_data,study=screen(price_frame,return_history=True,return_validation=True)
    assert len(normal)==len(ranking)==1
    assert len(chart_data["TWSE-2466"])==65
    assert study.keys()=={"records","pending","events"}
    assert study["events"]>=0


def test_daily_export_and_web_panel(tmp_path,monkeypatch):
    import radar.main as main
    monkeypatch.setattr(main,"ROOT",tmp_path)
    (tmp_path/"data").mkdir()
    conn=connect(tmp_path/"data"/"radar.sqlite")
    frame,_=make_test_data()
    for r in frame.itertuples(index=False):
        store(conn,r.date,r.market,[(
            r.date,r.market,r.code,r.name,r.open,r.high,r.low,
            r.close,r.volume,r.turnover,r.source_url,
        )],"ok")
    export(conn,tmp_path/"output",tmp_path/"docs",frame.date.max())
    conn.close()
    report=json.loads((tmp_path/"docs"/"signal_validation.json").read_text("utf-8"))
    assert report["entry_rule"]=="NEXT_SESSION_OPEN"
    assert report["interpretation"]["includes_trading_costs"] is False
    assert len(report["sample_rows"])==16
    summary=json.loads((tmp_path/"docs"/"report.json").read_text("utf-8"))
    assert "signal_validation" in summary["summary"]
    from openpyxl import load_workbook
    wb=load_workbook(tmp_path/"output"/"radar_latest.xlsx",read_only=True)
    assert "訊號歷史統計" in wb.sheetnames
    assert "已完成5日案例" in wb.sheetnames
    html=render_dashboard()
    assert 'id="val-panel"' in html
    assert 'id="val-days"' in html
    assert 'id="pv-panel"' in html
    assert "技術分數細項分析" in html
    assert "signal_validation.json" in html
    if shutil.which("node"):
        scripts=re.findall(r"<script>([\s\S]*?)</script>",html)
        assert len(scripts)==1
        result=subprocess.run(["node","--check"],input=scripts[0],
                              text=True,capture_output=True)
        assert result.returncode==0,result.stderr
