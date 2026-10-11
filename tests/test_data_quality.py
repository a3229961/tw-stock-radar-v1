"""Real-data safety checks for the V5 market completeness pipeline."""
import datetime as dt
import json

import pandas as pd
import pytest

from radar.data_quality import (
    MarketBatchRejected, audit_candidate, make_quality_report,
    MIN_BASE_ROWS,
)
from radar.sources import parse_twse, parse_tpex
from radar.storage import connect, store, record_attempt, replace_market_snapshot, read_prices
from radar.main import days_to_recheck, export
from radar.dashboard import render_dashboard


def rows(day, market, n, first=1000, close=25.0):
    return [(
        day, market, f"{first+i:04d}", f"測試{i}",
        close-.5, close+1, close-1, close,
        1000, 1000*close, "https://www.twse.com.tw/official-test"
    ) for i in range(n)]


def test_minimum_and_relative_market_width(tmp_path):
    conn=connect(tmp_path/"db.sqlite")
    date="2026-10-08"
    for market,n in [("TWSE",650),("TPEX",450)]:
        result=audit_candidate(conn,date,market,rows(date,market,n))
        assert result["row_count"]==n
    with pytest.raises(MarketBatchRejected,match="probable partial feed"):
        audit_candidate(conn,date,"TWSE",rows(date,"TWSE",450))
    with pytest.raises(MarketBatchRejected,match="probable partial feed"):
        audit_candidate(conn,date,"TPEX",rows(date,"TPEX",350))
    # Bootstrap is already validated; an established 1100-share market can't
    # silently shrink to 700 just because 700 exceeds the absolute minimum.
    for i in range(3):
        d=f"2026-10-0{i+1}"
        replace_market_snapshot(conn,d,"TWSE",rows(d,"TWSE",1100))
    with pytest.raises(MarketBatchRejected,match="historical median 1100"):
        audit_candidate(conn,date,"TWSE",rows(date,"TWSE",700))
    conn.close()


def test_bad_ohlc_duplicate_and_turnover_drift(tmp_path):
    conn=connect(tmp_path/"db.sqlite")
    good=rows("2026-10-08","TWSE",650)
    broken=good.copy()
    bad=list(broken[0])
    bad[5]=bad[6]-1
    broken[0]=tuple(bad)
    with pytest.raises(MarketBatchRejected,match="invalid values"):
        audit_candidate(conn,"2026-10-08","TWSE",broken)
    duplicated=good.copy()
    duplicated[1]=duplicated[0]
    with pytest.raises(MarketBatchRejected,match="duplicate"):
        audit_candidate(conn,"2026-10-08","TWSE",duplicated)
    reversed_value=good.copy()
    for i in range(100):
        r=list(reversed_value[i])
        r[9]=1000*1000   # above quoted high x 1000 shares
        reversed_value[i]=tuple(r)
    with pytest.raises(MarketBatchRejected,match="field mapping/units drift"):
        audit_candidate(conn,"2026-10-08","TWSE",reversed_value)
    conn.close()


def test_atomic_replace_updates_revisions_and_preserves_after_rejected_fetch(tmp_path):
    conn=connect(tmp_path/"db.sqlite")
    day="2026-10-08"
    original=rows(day,"TWSE",650)
    replace_market_snapshot(conn,day,"TWSE",original)
    revised=rows(day,"TWSE",650,first=1001)
    edited=list(revised[0])
    edited[7]=25.25
    revised[0]=tuple(edited)
    result=audit_candidate(conn,day,"TWSE",revised)
    assert result["row_count"]==650
    changed=replace_market_snapshot(conn,day,"TWSE",revised)
    assert changed=={"changed":1,"added":1,"removed":1,"rows":650}
    assert conn.execute(
        "SELECT COUNT(*) FROM prices WHERE date=? AND market=?",(day,"TWSE")
    ).fetchone()[0]==650
    assert conn.execute(
        "SELECT close FROM prices WHERE date=? AND market='TWSE' AND code='1001'",(day,)
    ).fetchone()[0]==25.25
    record_attempt(conn,day,"TWSE","quality_rejected",320,"partial official feed")
    assert conn.execute("SELECT status FROM fetch_log WHERE date=? AND market='TWSE'",(day,)).fetchone()[0]=="ok"
    assert conn.execute("SELECT status FROM fetch_attempts ORDER BY id DESC LIMIT 1").fetchone()[0]=="quality_rejected"
    # No partial snapshot was substituted.
    assert conn.execute("SELECT COUNT(*) FROM prices").fetchone()[0]==650
    conn.close()


def test_duplicate_source_quotes_are_rejected_without_silent_overwrite():
    day=dt.date(2026,10,8)
    fields=["證券代號","證券名稱","成交股數","成交金額","開盤價","最高價","最低價","收盤價"]
    a=["1234","甲", "1000","10000","10","11","9","10"]
    b=a.copy();b[7]="10.5"
    with pytest.raises(ValueError,match="Conflicting official source rows"):
        parse_twse({"stat":"OK","tables":[{"fields":fields,"data":[a,b]}]},day,
                   "https://www.twse.com.tw/")
    tp_a=["1234","甲","10","0","10","11","9","10","1000","10000"]
    tp_b=tp_a.copy();tp_b[2]="10.5"
    with pytest.raises(ValueError,match="Conflicting official source rows"):
        parse_tpex({"aaData":[tp_a,tp_b]},day,"https://www.tpex.org.tw/")


def test_quality_report_identifies_latest_joint_day_and_failed_refresh(tmp_path):
    conn=connect(tmp_path/"db.sqlite")
    day="2026-10-08"
    replace_market_snapshot(conn,day,"TWSE",rows(day,"TWSE",650))
    replace_market_snapshot(conn,day,"TPEX",rows(day,"TPEX",450,first=5000))
    record_attempt(conn,day,"TWSE","quality_rejected",10,"bad recheck")
    later="2026-10-09"
    store(conn,later,"TWSE",rows(later,"TWSE",650), "ok")
    prices=read_prices(conn)
    info=make_quality_report(conn,prices,dt.date(2026,10,11))
    assert info["last_both_markets_verified_date"]==day
    assert info["latest_verified_market_counts"]=={"TWSE":650,"TPEX":450}
    assert info["last_45d_problem_snapshots"]>=1
    assert info["session_coverage"]["status"]=="INSUFFICIENT_COMPLETE_SESSIONS"
    assert days_to_recheck(conn,1)==[dt.date(2026,10,8)]
    conn.close()


def test_export_stays_on_joint_market_day_and_renders_qa(tmp_path,monkeypatch):
    import radar.main as main
    monkeypatch.setattr(main,"ROOT",tmp_path)
    (tmp_path/"data").mkdir()
    conn=connect(tmp_path/"data"/"radar.sqlite")
    good="2026-10-08"
    store(conn,good,"TWSE",rows(good,"TWSE",1),"ok")
    store(conn,good,"TPEX",rows(good,"TPEX",1,first=5000),"ok")
    store(conn,"2026-10-09","TWSE",rows("2026-10-09","TWSE",1),"ok")
    export(conn,tmp_path/"output",tmp_path/"docs",dt.date(2026,10,11))
    conn.close()
    report=json.loads((tmp_path/"docs"/"report.json").read_text("utf-8"))
    qa=json.loads((tmp_path/"docs"/"data_quality.json").read_text("utf-8"))
    assert report["summary"]["latest_data_date"]==good
    assert qa["last_price_date_in_database"]=="2026-10-09"
    assert qa["last_both_markets_verified_date"]==good
    assert report["summary"]["data_quality"]["status"]=="PARTIAL_COVERAGE_NOT_CERTIFIED"
    page=(tmp_path/"docs"/"index.html").read_text("utf-8")
    assert 'id="qa-panel"' in page
    assert 'data_quality.json' in page
    from openpyxl import load_workbook
    xlsx=load_workbook(tmp_path/"output"/"radar_latest.xlsx",read_only=True)
    assert "資料品質總覽" in xlsx.sheetnames
    assert "待補抓與異常" in xlsx.sheetnames
