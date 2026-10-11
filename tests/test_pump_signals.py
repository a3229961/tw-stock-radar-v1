"""Offline verification for point-in-time OHLCV push signals and published artifacts."""
import json
import math
import re
import shutil
import subprocess

import pandas as pd

from radar.engine import indicators, screen
from radar.main import export
from radar.pump_export import publish_impulse
from radar.storage import connect, store
from radar.dashboard import render_dashboard


def candles(days=105, signal_at=94, turnover=28_000_000):
    dates = pd.bdate_range("2026-04-01", periods=days)
    rows = []
    for i, day in enumerate(dates):
        close = 100 + 2.5 * math.sin(i * .37)
        hi = close + 1
        lo = close - 1
        op = close - .3
        vol = 215000
        if i == signal_at:
            close = max(r["high"] for r in rows[-20:]) + 2.0
            op = close - 3
            hi = close + 1
            lo = op - .7
            vol = 850000
        if i > signal_at:
            close = rows[-1]["close"] + (1.5 if i%2==0 else 1)
            op = close-1.1
            hi = close+1
            lo = close-2
            vol = 390000
        rows.append({
            "date": day.strftime("%Y-%m-%d"), "market": "TWSE",
            "code": "2466", "name": "冠西電",
            "open": op, "high": hi, "low": lo, "close": close,
            "volume": vol, "turnover": turnover if i!=signal_at else 80_000_000,
            "source_url": "synthetic-test-only",
        })
    return pd.DataFrame(rows)


def test_price_volume_launch_and_no_future_leakage():
    data = candles(days=103, signal_at=94)
    before = indicators(data.iloc[:95])
    extended = indicators(data)
    previous = before.iloc[-1]
    actual = extended.iloc[94]
    assert previous.impulse_signal == actual.impulse_signal
    assert previous.impulse_strength == actual.impulse_strength
    assert previous.impulse_return5_pct == actual.impulse_return5_pct
    assert actual.impulse_signal in ("START", "EARLY")
    assert actual.impulse_event


def test_low_turnover_cannot_be_labeled_as_launch():
    data = candles(days=95, signal_at=94)
    data.loc[data.index[-1], "turnover"] = 400000
    result = indicators(data)
    assert result.iloc[-1].impulse_signal == "NONE"
    assert pd.isna(result.iloc[-1].impulse_strength)


def test_candlestick_shards_and_report_without_broker_data(tmp_path, monkeypatch):
    import radar.main as main
    monkeypatch.setattr(main, "ROOT", tmp_path)
    (tmp_path/"data").mkdir()
    con=connect(tmp_path/"data"/"radar.sqlite")
    df=candles(days=95,signal_at=94)
    for row in df.itertuples(index=False):
        store(con,row.date,row.market,[(
            row.date,row.market,row.code,row.name,row.open,row.high,row.low,
            row.close,row.volume,row.turnover,row.source_url
        )],"ok")
    export(con,tmp_path/"output",tmp_path/"docs",df.date.max())
    con.close()
    output=tmp_path/"docs"
    feed=json.loads((output/"impulse.json").read_text("utf-8"))
    assert feed["summary"]["latest_data_date"]==df.date.max()
    assert feed["summary"]["branch_verified"] is False
    assert feed["summary"]["new_start_count"]+feed["summary"]["new_early_count"]>=1
    assert feed["candidates"][0]["code"]=="2466"
    chart=json.loads((output/"charts"/"TWSE-2.json").read_text("utf-8"))
    bars=chart["stocks"]["2466"]
    assert len(bars)==65
    assert bars[-1][8] in ("START","EARLY")
    assert bars[-1][9] is True
    report=json.loads((output/"report.json").read_text("utf-8"))
    assert report["impulse_top"][0]["code"]=="2466"
    assert report["summary"]["impulse_signals"]["branch_verified"] is False
    excel=tmp_path/"output"/"radar_latest.xlsx"
    from openpyxl import load_workbook
    wb=load_workbook(excel,read_only=True)
    assert "量價起漲候選" in wb.sheetnames
    assert "短線過熱警示" in wb.sheetnames


def test_homepage_contains_signal_panel_and_valid_javascript():
    html=render_dashboard()
    assert 'id="pv-panel"' in html
    assert 'id="pv-chart"' in html
    assert 'id="pv-filter"' in html
    assert "impulse.json" in html
    assert "charts/" in html
    assert "技術分數細項分析" in html
    if shutil.which("node"):
        js=re.findall(r"<script>([\s\S]*?)</script>",html)
        assert len(js)==1
        check=subprocess.run(["node","--check"],input=js[0],text=True,capture_output=True)
        assert check.returncode==0,check.stderr


def test_last_day_has_no_signal_without_80_bars():
    result=indicators(candles(days=79,signal_at=78))
    assert result.iloc[-1].impulse_signal == "NONE"
